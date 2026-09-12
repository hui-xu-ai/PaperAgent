# -*- coding: utf-8 -*-
"""JCR 分区表 xlsx 导入（openpyxl）。

- 后台解析（导入/预览**只回摘要**，不 dump 全文——用户要求省 token）
- 识别两个 sheet：2025JCRIF-分区（jcr 表）与 2025中科学院分区表（cas 表），
  年份从 sheet 名提取（YYYY）
- 预览 → 确认 → upsert journals.db（用户重新提供表格 = 更新）
"""
from __future__ import annotations

import re
from pathlib import Path

_JCR_HEADER_MAP = {
    "期刊名称": "journal_name",
    "2024JIF": "jif",
    "quartile": "quartile",
    "jif rank": "jif_rank",
    "2023分区": "zone_2023",
    "total citation": "total_citation",
    "category": "category",
    "issn": "issn",
    "eissn": "eissn",
}
_CAS_HEADER_MAP = {
    "期刊名称": "journal_name",
    "2025分区": "zone",
    "top": "is_top",
    "open access": "is_oa",
}


def _norm_hdr(h: str) -> str:
    return re.sub(r"\s+", " ", (h or "").strip()).lower()


def _sheet_year(sheet_name: str) -> int:
    m = re.search(r"(20\d{2})", sheet_name or "")
    return int(m.group(1)) if m else 0


def _read_rows(ws) -> tuple[list[str] | None, list[list]]:
    """读 sheet：返回 (header, rows)。自适应跳过标题行（真实中科院表首行是标题）。

    判据：首行匹配 jcr/cas 列模式即视为 header；否则视作标题行继续读下一行。
    """
    it = ws.iter_rows(values_only=True)
    header: list | None = None
    for _ in range(3):
        row = next(it, None)
        if row is None:
            return None, []
        cand = list(row)
        if _is_jcr_sheet(cand) or _is_cas_sheet(cand):
            header = cand
            break
        header = cand  # 标题行（继续读下一行）
    rows = [list(r) for r in it]
    return header, rows


def _parse_jcr(header: list, rows: list[list], year: int) -> list[dict]:
    idx = {_norm_hdr(h): i for i, h in enumerate(header) if h}
    out: list[dict] = []
    for r in rows:
        d = {"year": year}
        for en, cn in _JCR_HEADER_MAP.items():
            i = idx.get(_norm_hdr(en))
            if i is not None and i < len(r) and r[i] is not None:
                d[cn] = str(r[i]).strip()
        if d.get("journal_name"):
            out.append(d)
    return out


def _parse_cas(header: list, rows: list[list], year: int) -> list[dict]:
    idx = {_norm_hdr(h): i for i, h in enumerate(header) if h}
    out: list[dict] = []
    for r in rows:
        d = {"year": year}
        for en, cn in _CAS_HEADER_MAP.items():
            i = idx.get(_norm_hdr(en))
            if i is not None and i < len(r) and r[i] is not None:
                d[cn] = str(r[i]).strip()
        if d.get("journal_name"):
            out.append(d)
    return out


def _is_jcr_sheet(header: list) -> bool:
    """JCR sheet 判据：含 quartile 或 JIF 列（'期刊名称' 两表都有，不能作判据）。"""
    h = {_norm_hdr(x) for x in header if x}
    return "quartile" in h or "jif" in h or any(
        _norm_hdr(x).startswith("20") and "jif" in _norm_hdr(x) for x in header if x)


def _is_cas_sheet(header: list) -> bool:
    """中科院 sheet 判据：含 Top 或 Open Access 列（中科院表特有；'YYYY分区' JCR 表也有）。"""
    h = {_norm_hdr(x) for x in header if x}
    return "top" in h or "open access" in h


def parse_xlsx(path: str | Path) -> dict:
    """解析 xlsx → {sheets: [{sheet, year, header, rows, jcr_count, cas_count}]}。

    只回摘要（每 sheet 行数/字段/样例 2 行），不 dump 全表内容。
    """
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets = []
    total_jcr = total_cas = 0
    for sn in wb.sheetnames:
        ws = wb[sn]
        header, rows = _read_rows(ws)
        if not header:
            continue
        year = _sheet_year(sn)
        jcr_rows = _parse_jcr(header, rows, year) if _is_jcr_sheet(header) else []
        cas_rows = _parse_cas(header, rows, year) if _is_cas_sheet(header) else []
        total_jcr += len(jcr_rows)
        total_cas += len(cas_rows)
        sheets.append({
            "sheet": sn, "year": year,
            "header": [h for h in header if h],
            "rows": len(rows),
            "jcr_count": len(jcr_rows), "cas_count": len(cas_rows),
            "sample": [dict(r) for r in (jcr_rows or cas_rows)[:2]],
        })
    wb.close()
    return {"sheets": sheets, "total_jcr": total_jcr, "total_cas": total_cas}


def import_xlsx(path: str | Path, journals_db) -> dict:
    """确认导入：解析 xlsx → upsert journals.db。返回导入统计摘要。"""
    parsed = parse_xlsx(path)
    jcr_rows, cas_rows = [], []
    for s in parsed["sheets"]:
        import openpyxl

        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb[s["sheet"]]
        header, rows = _read_rows(ws)
        year = s["year"]
        if _is_jcr_sheet(header):
            jcr_rows.extend(_parse_jcr(header, rows, year))
        if _is_cas_sheet(header):
            cas_rows.extend(_parse_cas(header, rows, year))
        wb.close()
    jcr_n = journals_db.upsert_jcr(jcr_rows)
    cas_n = journals_db.upsert_cas(cas_rows)
    return {"jcr_imported": jcr_n, "cas_imported": cas_n,
            "jcr_total": parsed["total_jcr"], "cas_total": parsed["total_cas"],
            "stats": journals_db.stats()}
